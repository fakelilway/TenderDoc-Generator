"""角色业绩勾选(项目经理/总工名下业绩人工多选)的服务层测试。

覆盖:候选按选派人名下取、勾选存对应列、None/[]语义、换人自动清空勾选。
"""
from typing import Any

from services import project_service
from services.project import performance as perf_mod
from services.project import personnel as pers_mod


def _record(i: int, manager: str = "李刚", tech: str = "") -> dict:
    return {
        "project_name": f"测试业绩{i}号工程",
        "project_type": "公路工程",
        "project_year": 2020 + i,
        "amount_wan": 100.0 * i,
        "contract_price": f"{i}000000元",
        "project_manager": manager,
        "tech_leader": tech,
        "document_id": i,
    }


class _Cur:
    """假游标:按序吐 fetchone 结果,记录 execute 的 SQL/参数。"""

    def __init__(self, rows: list[Any]):
        self._rows = list(rows)
        self.executed: list[tuple[str, Any]] = []

    def execute(self, sql: str, params: Any = None) -> None:
        self.executed.append((sql, params))

    def fetchone(self) -> Any:
        return self._rows.pop(0) if self._rows else None

    def __enter__(self) -> "_Cur":
        return self

    def __exit__(self, *args: Any) -> bool:
        return False


class _Conn:
    def __init__(self, cur: _Cur):
        self._cur = cur

    def cursor(self, cursor_factory: Any = None) -> _Cur:
        return self._cur

    def __enter__(self) -> "_Conn":
        return self

    def __exit__(self, *args: Any) -> bool:
        return False


def test_recommend_role_performance_lists_person_records(monkeypatch) -> None:
    """选派了总工许明英 → 候选=她名下(tech_leader)记录,带证明标记;selected 原样透传。"""
    monkeypatch.setattr(
        perf_mod,
        "_fetch_project",
        lambda pid: {
            "selected_personnel": {"tech_director": {"name": "许明英"}},
            "selected_td_performance": None,  # 没勾过
        },
    )
    from services import similar_project_info_service as spi

    monkeypatch.setattr(
        spi, "records_for_tech_leader",
        lambda name: [_record(1, tech=name), _record(2, tech=name)] if name == "许明英" else [],
    )
    from services import performance_archive_service as pa

    monkeypatch.setattr(pa, "list_evidence_groups", lambda: {"测试业绩1号工程": []})

    res = perf_mod.recommend_role_performance(7, "td")
    assert res["person"] == "许明英"
    assert res["selected"] is None  # 没勾过(全部人工手选,前端一条不预勾)
    names = [r["name"] for r in res["recommendations"]]
    assert names == ["测试业绩1号工程", "测试业绩2号工程"]
    assert res["recommendations"][0]["has_evidence"] is True
    assert res["recommendations"][0]["matched"] == ["有证明扫描"]
    assert res["recommendations"][1]["gaps"] == ["缺证明扫描"]


def test_recommend_role_performance_without_person(monkeypatch) -> None:
    """该角色尚未选派 → person=None、候选为空(前端提示先选派)。"""
    monkeypatch.setattr(
        perf_mod, "_fetch_project", lambda pid: {"selected_personnel": {}}
    )
    res = perf_mod.recommend_role_performance(7, "pm")
    assert res["person"] is None
    assert res["recommendations"] == []


def test_recommend_role_performance_lists_role_holders(monkeypatch) -> None:
    """响应带 role_holders(信息表里当过该角色的人+条数):选了47表外的人时,
    面板要能告诉用户"谁有业绩可选"(凌雨实测困惑)。"""
    monkeypatch.setattr(
        perf_mod,
        "_fetch_project",
        lambda pid: {"selected_personnel": {"tech_director": {"name": "凌雨"}}},
    )
    from services import similar_project_info_service as spi

    monkeypatch.setattr(spi, "records_for_tech_leader", lambda name: [])
    monkeypatch.setattr(
        spi,
        "list_similar_project_records",
        lambda: [
            {"project_name": f"工程{i}", "tech_leader": "许明英", "project_manager": "李刚"}
            for i in range(3)
        ] + [{"project_name": "工程x", "tech_leader": "赵勇", "project_manager": "李刚"}],
    )
    res = perf_mod.recommend_role_performance(7, "td")
    assert res["person"] == "凌雨"
    assert res["recommendations"] == []
    assert res["role_holders"][0] == {"name": "许明英", "count": 3}
    assert {"name": "赵勇", "count": 1} in res["role_holders"]


def test_save_selected_role_performance_writes_role_column(monkeypatch) -> None:
    """保存经理勾选 → 写 selected_pm_performance 列,条目清洗成5键。"""
    cur = _Cur([{"id": 7, "selected_pm_performance": [{"name": "测试业绩1号工程"}]}])
    monkeypatch.setattr(project_service, "_connect", lambda: _Conn(cur))

    res = perf_mod.save_selected_role_performance(
        7, "pm", [{"name": "测试业绩1号工程", "year": "2021", "extra": "丢弃"}, {"name": ""}]
    )
    sql, params = cur.executed[0]
    assert "selected_pm_performance" in sql
    clean = params[0].adapted  # psycopg2 Json 包装的原值
    assert clean == [
        {"name": "测试业绩1号工程", "year": "2021", "amount": "", "type": "", "document_id": None}
    ]
    assert res["role"] == "pm"
    assert res["selected"] == [{"name": "测试业绩1号工程"}]


def test_save_selected_role_performance_rejects_bad_role() -> None:
    import pytest

    with pytest.raises(ValueError):
        perf_mod.save_selected_role_performance(7, "boss", [])


def test_changing_person_resets_role_performance(monkeypatch) -> None:
    """换项目经理 → selected_pm_performance 一并重置为 NULL;同人重存则不动勾选。"""
    # 换人:李刚 → 王露
    cur = _Cur([
        {"selected_personnel": {"project_manager": {"name": "李刚"}}},
        {"id": 7, "selected_personnel": {"project_manager": {"name": "王露"}}},
    ])
    monkeypatch.setattr(project_service, "_connect", lambda: _Conn(cur))
    pers_mod.save_selected_project_manager(7, {"name": "王露"})
    update_sql = cur.executed[1][0]
    assert "selected_pm_performance = NULL" in update_sql

    # 同人重存:不清勾选
    cur2 = _Cur([
        {"selected_personnel": {"project_manager": {"name": "王露"}}},
        {"id": 7, "selected_personnel": {"project_manager": {"name": "王露"}}},
    ])
    monkeypatch.setattr(project_service, "_connect", lambda: _Conn(cur2))
    pers_mod.save_selected_project_manager(7, {"name": "王露"})
    assert "selected_pm_performance" not in cur2.executed[1][0]

    # 清空选派:也要清勾选
    cur3 = _Cur([
        {"selected_personnel": {"tech_director": {"name": "许明英"}}},
        {"id": 7, "selected_personnel": {}},
    ])
    monkeypatch.setattr(project_service, "_connect", lambda: _Conn(cur3))
    pers_mod.save_selected_tech_director(7, None)
    assert "selected_td_performance = NULL" in cur3.executed[1][0]


def test_llm_purpose_routing_tech_vs_default():
    """三路分流:商务卷走全局(kimi),技术卷走 TECH_LLM_PROVIDER(deepseek),
    解析走 PARSER_LLM_PROVIDER(deepseek)——2026-07-16/07-29 两次拍板。"""
    from types import SimpleNamespace

    from core.llm_client import resolve_llm_config

    s = SimpleNamespace(
        bid_llm_provider="kimi",
        tech_llm_provider="deepseek",
        parser_llm_provider="deepseek",
        kimi_api_key="sk-kimi-real-key-123456",
        kimi_base_url="https://api.moonshot.cn/v1",
        kimi_model="kimi-k3",
        deepseek_api_key="sk-ds-real-key-123456",
        deepseek_base_url="https://api.deepseek.com",
        deepseek_model="deepseek-v4-pro",
    )
    assert resolve_llm_config(s)[2] == "kimi-k3"
    assert resolve_llm_config(s, purpose="technical")[2] == "deepseek-v4-pro"
    assert resolve_llm_config(s, purpose="parser")[2] == "deepseek-v4-pro"
    # 没配覆盖 → 各用途都跟随全局
    s.tech_llm_provider = ""
    s.parser_llm_provider = ""
    assert resolve_llm_config(s, purpose="technical")[2] == "kimi-k3"
    assert resolve_llm_config(s, purpose="parser")[2] == "kimi-k3"


def test_chat_completion_strips_temperature_for_kimi():
    """kimi-k3 只接受 temperature=1:kimi 系剥掉非1温度,别家原样透传(实测400修复)。"""
    from types import SimpleNamespace

    from core.llm_client import chat_completion

    captured = {}

    class _FakeClient:
        class chat:
            class completions:
                @staticmethod
                def create(**kw):
                    captured.update(kw)
                    return SimpleNamespace(choices=[])

    chat_completion(_FakeClient, model="kimi-k3", messages=[], temperature=0)
    assert "temperature" not in captured
    captured.clear()
    chat_completion(_FakeClient, model="deepseek-v4-pro", messages=[], temperature=0)
    assert captured.get("temperature") == 0


def test_curated_resume_takes_priority_over_roster(monkeypatch) -> None:
    """资历表定稿(用户2026-07-31提供)优先于台账/OCR;拟任职务仍按选派角色。"""
    from services import v2_generation_service as v2

    import services.curated_resume_service as crs
    monkeypatch.setattr(
        crs, "get_curated_resume",
        lambda name: {
            "年龄": "50", "职称": "高级工程师", "学历": "专科",
            "工作年限": "22年", "毕业学校": "2018年1月毕业于国家开放大学",
            "拟任职务": "项目总工",  # 文档里写的职务,必须被忽略
        } if name == "许明英" else {},
    )
    import services.personnel_roster_service as prs
    monkeypatch.setattr(
        prs, "get_personnel_roster",
        lambda: {"roster": [{"name": "许明英", "title": "工程师",
                             "id_number": "34010119760101002X"}]},
    )
    f = v2.build_pm_resume_fields("许明英", role="项目经理")
    assert f["职称"] == "高级工程师"  # 定稿赢台账("工程师")
    assert f["年龄"] == "50"          # 定稿赢身份证推算
    assert f["学历"] == "专科" and f["工作年限"] == "22年"
    assert f["拟任职务"] == "项目经理"  # 按选派角色,不照抄文档
    assert f["性别"] == "女"           # 定稿没有的,台账补


def test_resume_parser_extracts_fields_from_doc_table() -> None:
    """导入脚本的表格解析:标签右邻取值、合并格去重、年龄取数字。"""
    from docx import Document
    sys_path_hack = None
    import importlib.util, pathlib
    spec = importlib.util.spec_from_file_location(
        "import_candidate_resumes",
        pathlib.Path(__file__).resolve().parents[1] / "scripts/import_candidate_resumes.py",
    )
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)

    doc = Document()
    t = doc.add_table(rows=3, cols=6)
    for i, v in enumerate(("姓 名", "李刚", "年 龄", "37岁", "专业", "公路工程")):
        t.cell(0, i).text = v
    for i, v in enumerate(("技术职称", "工程师", "学历", "专科", "拟在本标段工程任职", "项目经理")):
        t.cell(1, i).text = v
    for i, v in enumerate(("工作年限", "15年", "类似施工经验年限", "15年", "获奖情况", "无")):
        t.cell(2, i).text = v

    f = mod.parse_resume_table(t)
    assert f["姓名"] == "李刚" and f["年龄"] == "37"
    assert f["职称"] == "工程师" and f["工作年限"] == "15年"
    assert f["类似施工经验年限"] == "15年"
    assert "拟任职务" not in f  # 职务不入库,按选派定
