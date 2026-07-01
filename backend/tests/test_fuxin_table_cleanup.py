"""福昕转换后处理:理顺被切开的两字标签("性 别"→"性别"),不碰值。"""

from docx import Document


def test_normalize_split_labels_collapses_known_labels_only() -> None:
    from services.original_docx_format_service import _normalize_split_labels

    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "性  别"  # 已知标签被切开 → 理顺
    t.cell(0, 1).text = "电 话"
    t.cell(1, 0).text = "许 明 英"  # 值(非白名单标签)→ 不动
    t.cell(1, 1).text = "技术职称"  # 无空格 → 不动
    n = _normalize_split_labels(doc)
    assert t.cell(0, 0).text == "性别"
    assert t.cell(0, 1).text == "电话"
    assert t.cell(1, 0).text == "许 明 英"  # 值不碰
    assert n == 2
