"""商务卷"近年完成的类似项目"六种节的按人自动填表(信息表+情况表)。

招标商务格式里这套表通常成对出现:每节先一张**情况表**(汇总:业绩序号|项目名称|备注),
再跟"一业绩一张"的**信息表**(竖表:项目名称/发包人三要素/合同价格/开工交工日期/项目
经理…)。共六种节:{投标人, 项目经理, 项目总工} × {资格审查, 详细评审}。

业务规则(用户拍板,见 memory selection-features / import_similar_project_info):
  · 投标人(公司)节 → **业绩跟用户选的走**:照 selected_performance(业绩面板勾选的项目)填,
    逐条从全部《类似项目信息表》记录按项目名找到、原样补全字段;选经理不影响这张表;
  · 项目经理节 → 填选定项目经理名下的项目(跟人走);
  · 项目总工节 → 按记录里"技术负责人"字段归属,匹配不到就留白(绝不硬凑);
  · 记录多于模板空表 → 克隆整套(标题段+表+注段)补齐;少于 → 裁掉多余空表;
  · 汇总情况表行数不够 → 克隆空行补齐。

数据流:业绩面板存 projects.selected_performance;选派存 projects.selected_personnel →
workflow_service._apply_selected_project_manager 注入 profile["selected_performance"(选中业绩)/
"similar_projects_all"(全部信息表记录,供投标人节按名补全)/"similar_projects_pm"/
"similar_projects_td"](来自 similar_project_info_service)→ cloud_pdf_convert 在通用填表**之前**调本模块。
本模块处理过/认领过的表会进返回值 handled_tables,通用 _fill_performance_table 按它跳过,
免得把项目经理的业绩名再填进总工节的空汇总表。
"""

from __future__ import annotations

import logging
import re
from copy import deepcopy
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from services.original_docx_format_service import (
    _is_blank_or_placeholder,
    _is_similar_project_detail_table,
    _set_cell_value,
)

logger = logging.getLogger(__name__)

# 详细信息表行标签 → 记录字段(similar_project_info_service.RECORD_FIELDS)。
# 最长匹配;招标措辞各家略有出入,给常见同义写法留了别名。
_DETAIL_LABEL_MAP: tuple[tuple[str, str], ...] = (
    ("项目名称", "project_name"),
    ("合同名称", "project_name"),
    ("工程名称", "project_name"),
    ("项目所在地", "location"),
    ("工程所在地", "location"),
    ("发包人名称", "owner_name"),
    ("发包人地址", "owner_address"),
    ("发包人电话", "owner_phone"),
    ("发包人联系电话", "owner_phone"),
    ("合同价格", "contract_price"),
    ("合同金额", "contract_price"),
    ("中标价格", "contract_price"),
    ("开工日期", "start_date"),
    ("交工日期", "end_date"),
    ("竣工日期", "end_date"),
    ("完工日期", "end_date"),
    ("承担的工作", "work_scope"),
    ("工程质量", "quality"),
    ("项目经理", "project_manager"),
    ("注册建造师", "project_manager"),
    ("项目总工", "tech_leader"),
    ("技术负责人", "tech_leader"),
    ("监理单位", "supervisor_org"),
    ("总监理工程师", "chief_supervisor"),
    ("项目描述", "description"),
)
# 这些标签留模板原样(备注列招标常预印"资格审查业绩";序号是台账内部编号,不外泄)
_DETAIL_SKIP_LABELS = ("备注", "序号")


def _norm_text(s: str) -> str:
    return re.sub(r"[\s　]+", "", s or "")


def _role_of_title(text: str) -> str | None:
    """段落是否切换"类似项目"节的归属人。返回 bidder/pm/td 或 None(不切换)。

    命中条件:含"近年完成的类似"或"类似项目情况表/信息表"字样。归属判定:
    项目总工(且无项目经理字样)→td;含项目经理→pm(涵盖"（六）项目经理（项目总工）…"
    父标题,子节标题随后会再切);含投标人→bidder;**无任何人称前缀→也是 bidder**——
    招标惯例里"（四）近年完成的类似项目情况表"这种裸标题就是问投标人(公司)的
    (埇桥实测:裸标题识别不出→整表没人认领→该填的不填)。
    """
    t = _norm_text(text)
    if not t:
        return None
    if ("近年完成的类似" not in t) and ("类似项目情况表" not in t) and ("类似项目信息表" not in t):
        return None
    has_pm = "项目经理" in t
    has_td = "项目总工" in t
    if has_td and not has_pm:
        return "td"
    if has_pm:
        return "pm"
    return "bidder"


_QUALIFIER_RE = re.compile(r"(资格审查|详细评审|评分)")


def _qualifier_of_title(text: str) -> str:
    """标题里的评审限定词(资格审查/详细评审/评分),没有返回空串。

    同一归属人常有资格审查+详细评审两套节:分组必须把限定词算进节键,
    否则相邻同归属节会被并成一组、业绩被拆开两节各填一半(对抗验证实测)。
    """
    m = _QUALIFIER_RE.search(_norm_text(text))
    return m.group(1) if m else ""


def _is_summary_table(table: Table) -> bool:
    """汇总"情况表":表头含 业绩序号 + 项目名称/合同名称/工程名称。"""
    try:
        if len(table.rows) < 2 or len(table.columns) < 2:
            return False
        headers = [c.text.strip() for c in table.rows[0].cells]
    except Exception:
        return False
    if not any("业绩序号" in h for h in headers):
        return False
    return any(("项目名称" in h or "合同名称" in h or "工程名称" in h) for h in headers)


def _page_break_paragraph() -> Any:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def _ensure_page_break_before(p_el: Any) -> None:
    """给段落设"段前分页"。比插入独立分页空段干净:福昕的标题常挂浮动文本框,
    多一个空段会跟浮框定位叠加挤出整页空白(实测)。

    按 CT_PPr 的子元素顺序插(pStyle/keepNext/keepLines 之后),不能无脑 append
    到 pPr 末尾——顺序违规的 docx 严格解析器会拒开。
    """
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p_el.insert(0, ppr)
    if ppr.find(qn("w:pageBreakBefore")) is not None:
        return
    pbb = OxmlElement("w:pageBreakBefore")
    anchor = None
    for tag in ("w:pStyle", "w:keepNext", "w:keepLines"):
        found = ppr.find(qn(tag))
        if found is not None:
            anchor = found
    if anchor is not None:
        anchor.addnext(pbb)
    else:
        ppr.insert(0, pbb)


def _strip_sectpr(el: Any) -> None:
    """克隆件里绝不能带 w:sectPr(节属性):复制节界会改后文页面设置。"""
    for sect in el.findall(f".//{qn('w:sectPr')}"):
        sect.getparent().remove(sect)


def _detail_table_is_blank(tbl_el: Any, document: Any) -> bool:
    """信息表的值格是否全空(备注/序号行除外)。裁剪多余表前必查:预印了内容的表不删。"""
    table = Table(tbl_el, document._body)
    for row in table.rows:
        cells = row.cells
        n = len(cells)
        i = 0
        while i < n:
            field = _label_field(cells[i].text)
            if field is None:
                i += 1
                continue
            j = i + 1
            while j < n and cells[j]._tc is cells[i]._tc:
                j += 1
            if j >= n:
                break
            if not _is_blank_or_placeholder(cells[j].text.strip()):
                return False
            i = j + 1
            while i < n and cells[i]._tc is cells[j]._tc:
                i += 1
    return True


def _is_padding_paragraph(el: Any) -> bool:
    """纯排版填充空段:无文字、无图、无换页、无节属性。"""
    if el.tag != qn("w:p"):
        return False
    if "".join(el.itertext()).strip():
        return False
    if el.find(f".//{qn('w:drawing')}") is not None:
        return False
    if el.find(f".//{qn('w:br')}") is not None:
        return False
    if el.find(f".//{qn('w:sectPr')}") is not None:
        return False
    return True


def _absorb_page_padding(last_el: Any, min_run: int = 3, max_scan: int = 18) -> int:
    """吸收表格区之后的"整页填充空段",改设下一段"段前分页"。

    福昕照抄的模板是"一页一版":页尾用一串空段垫满,让下一节从新页开始。表格被
    克隆/扩行撑长后,这些垫页空段会被挤到下一页变成整页空白(实测5处)。把它们
    删掉、给下一个非空段设段前分页,版面语义不变(下一节仍从新页开始)、空白页消失。

    垫页空段和表格之间可能隔着表单尾巴(注:/投标人(盖章)/日期,实测详细评审节),
    故向后扫描 max_scan 个元素找 ≥min_run 的连续空段;碰到下一张表或下一个
    "类似项目"节标题就停手(那是别人的地盘)。

    福昕的表单尾巴里常有"换栏符"段落(单栏页面=换页符):若扫描途中越过了它,
    垫页空段删掉即可、**不再**设段前分页——否则双重分页又挤出一张空白页(实测)。
    """
    el = last_el.getnext()
    scanned = 0
    crossed_col_break = False
    while el is not None and scanned < max_scan:
        if el.tag == qn("w:tbl"):
            return 0  # 到下一张表,不越界
        if el.tag == qn("w:p"):
            text = "".join(el.itertext()).strip()
            if text and _role_of_title(text):
                return 0  # 已到下一节标题
            if any(
                br.get(qn("w:type")) in ("column", "page")
                for br in el.findall(f".//{qn('w:br')}")
            ):
                crossed_col_break = True
            if _is_padding_paragraph(el):
                run: list[Any] = [el]
                probe = el.getnext()
                while probe is not None and _is_padding_paragraph(probe):
                    run.append(probe)
                    probe = probe.getnext()
                if len(run) >= min_run and probe is not None and probe.tag == qn("w:p"):
                    for empty in run:
                        empty.getparent().remove(empty)
                    if not crossed_col_break:
                        _ensure_page_break_before(probe)
                    return len(run)
                scanned += len(run)
                el = probe
                continue
        el = el.getnext()
        scanned += 1
    return 0


def _adjacent_title_p(tbl_el: Any) -> Any | None:
    """表格紧前的标题段(跳过空段,最多回看4个元素;必须像"类似项目…表"标题)。"""
    el = tbl_el.getprevious()
    for _ in range(4):
        if el is None or not el.tag.endswith("}p"):
            return None
        text = "".join(el.itertext()).strip()
        if text:
            return el if _role_of_title(text) else None
        el = el.getprevious()
    return None


_NOTE_HEAD_RE = re.compile(r"^[注註]\s*[：:1１]")


def _adjacent_note_p(tbl_el: Any) -> Any | None:
    """表格紧后的"注：…"段(跳过空段,最多前看4个元素)。

    只认"注：/注:/注1"开头,不认裸"注"字打头——"注册建造师""注意事项"这类
    标题绝不能被当注段克隆/删除。
    """
    el = tbl_el.getnext()
    for _ in range(4):
        if el is None or not el.tag.endswith("}p"):
            return None
        text = "".join(el.itertext()).strip()
        if text:
            return el if _NOTE_HEAD_RE.match(text) else None
        el = el.getnext()
    return None


def _label_field(text: str) -> str | None:
    """单元格文本 → 记录字段名;不是可填标签返回 None。标签必须短(真标签≤20字)。"""
    label = _norm_text(text).replace("（", "(").replace("）", ")")
    if not label or len(label) > 20:
        return None
    if any(label.startswith(k) for k in _DETAIL_SKIP_LABELS):
        return None
    matches = [(k, f) for k, f in _DETAIL_LABEL_MAP if k in label]
    if not matches:
        return None
    return max(matches, key=lambda kf: len(kf[0]))[1]


def _fill_detail_table(table: Table, record: dict[str, Any]) -> int:
    """按标签把一条记录原样填进一张竖表。只填空格(保真)。返回填入格数。

    逐格行走而非只看首列:兼容"开工日期|值|交工日期|值"这种一行两对字段的排版,
    每个标签格的值写进它右侧第一个不同单元格,绝不写进下一对标签的值格。
    """
    filled = 0
    for row in table.rows:
        cells = row.cells
        n = len(cells)
        i = 0
        while i < n:
            field = _label_field(cells[i].text)
            if field is None:
                i += 1
                continue
            # 值格 = 本行右侧第一个不同 tc 的格(合并单元格会重复吐出同一 tc)
            j = i + 1
            while j < n and cells[j]._tc is cells[i]._tc:
                j += 1
            if j >= n:
                break  # 整行合并成一个格(节标题行之类),没有值格
            value = str(record.get(field) or "").strip()
            if value and _is_blank_or_placeholder(cells[j].text.strip()):
                _set_cell_value(cells[j], value)
                filled += 1
            # 跳过值格(含其合并重复),同行可能还有下一对标签
            i = j + 1
            while i < n and cells[i]._tc is cells[j]._tc:
                i += 1
    return filled


def _clear_row_texts(tr_el: Any) -> None:
    """清空克隆出来的行的所有单元格文本(每格保留首段,删多余段/所有 run)。"""
    for tc in tr_el.findall(qn("w:tc")):
        paragraphs = tc.findall(qn("w:p"))
        for extra in paragraphs[1:]:
            tc.remove(extra)
        if paragraphs:
            for run in paragraphs[0].findall(qn("w:r")):
                paragraphs[0].remove(run)


def _fill_summary_table(table: Table, names: list[str]) -> int:
    """填汇总情况表:业绩序号列 1..N + 项目名称列。行不够克隆末行补齐。返回填入行数。"""
    if not names:
        return 0
    headers = [c.text.strip() for c in table.rows[0].cells]
    col_name = next(
        (c for c, h in enumerate(headers)
         if "项目名称" in h or "合同名称" in h or "工程名称" in h),
        None,
    )
    col_seq = next((c for c, h in enumerate(headers) if "序号" in h), None)
    if col_name is None:
        return 0
    # 可填行 = 项目名称列为空/占位的行;已有内容的行(招标示例等)不动
    fillable = [
        ri for ri in range(1, len(table.rows))
        if _is_blank_or_placeholder(
            table.cell(ri, col_name).text.strip().strip("…．.")
        )
        or table.cell(ri, col_name).text.strip() in ("……", "…", "...", "．．．")
    ]
    # 行数不够 → 克隆最后一个**可填空行**补齐(插在它后面,不追到表尾——
    # 表尾可能是跨列备注/脚注行,克隆它会把业绩名填进合并格)
    while len(fillable) < len(names) and fillable:
        base_idx = fillable[-1]
        base_tr = table.rows[base_idx]._tr
        new_tr = deepcopy(base_tr)
        _clear_row_texts(new_tr)
        base_tr.addnext(new_tr)
        fillable.append(base_idx + 1)
    filled = 0
    for i, name in enumerate(names):
        if i >= len(fillable):
            break  # 表连克隆都补不了(异常结构),静默截断前有日志
        ri = fillable[i]
        _set_cell_value(table.cell(ri, col_name), name)
        if col_seq is not None:
            seq_cell = table.cell(ri, col_seq)
            if _is_blank_or_placeholder(seq_cell.text.strip()) or seq_cell.text.strip() in ("……", "…", "..."):
                _set_cell_value(seq_cell, str(i + 1))
        filled += 1
    if filled < len(names):
        logger.warning("类似项目汇总表行数不足且无法扩行,%d 条业绩只填了 %d 行", len(names), filled)
    return filled


def _pick_by_selection(
    pool: list[dict[str, Any]], selected: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """按勾选列表(每条含 name)从记录池里挑出全字段记录,顺序跟勾选走。

    先按原名精确匹配(勾选项的名字本来就从记录原样带出,永远命中),匹配不上才退归一化
    模糊匹配——_norm 会剃掉括号补充说明,"(一标段)/(二标段)"这类两条不同业绩归一化后
    同名,直接用它建索引会把两条都填成同一条的数据。池里都找不到的只填项目名、其余留白
    待人工(不硬凑别的项目的数据)。"""
    from services.performance_archive_service import _norm

    by_exact: dict[str, dict[str, Any]] = {}
    by_norm: dict[str, dict[str, Any]] = {}
    for r in pool:
        raw = str(r.get("project_name") or "").strip()
        by_exact.setdefault(raw, r)
        by_norm.setdefault(_norm(raw), r)
    out: list[dict[str, Any]] = []
    for item in selected:
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        out.append(
            by_exact.get(name) or by_norm.get(_norm(name)) or {"project_name": name}
        )
    return out


def _records_for_role(role: str, profile: dict[str, Any]) -> list[dict[str, Any]]:
    """节归属人 → 该填的记录列表(顺序即出表顺序)。

    投标人(公司)节:**业绩跟用户选的走**——以 selected_performance(业绩面板勾选的项目)
    为准,逐条去全部信息表记录里按项目名找到它、原样补全字段(发包人三要素/合同价/监理/
    工期…);选的项目信息表里没有则只填项目名、其余留白待人工。没做业绩选择(键缺失)则
    留白不动模板,交人工填。选项目经理不影响这张表。
    项目经理/项目总工节:候选=选定人名下的项目;**全部人工手选**(用户2026-07-11拍板,
    不做默认全选):勾中哪几条填哪几条、顺序跟勾选走;没勾(None/键缺失/[])一律留白,
    与投标人节同一规矩。匹配不到留白,绝不拿别人的凑。
    """
    if role in ("td", "pm"):
        pool = list(profile.get(f"similar_projects_{role}") or [])
        selected = profile.get(f"selected_{role}_performance")
        return _pick_by_selection(pool, selected or [])
    # bidder(投标人/公司):照用户选中的业绩填,从全部信息表记录按名补全字段
    selected = profile.get("selected_performance")
    if selected is None:
        return []  # 没选业绩 → 留白待人工(主循环保留空模板,不乱填)
    pool = profile.get("similar_projects_all") or profile.get("similar_projects_pm") or []
    return _pick_by_selection(pool, selected)


def fill_similar_project_sections(document: Any, profile: dict[str, Any]) -> dict[str, Any]:
    """主入口:按节归属人填 情况表(汇总)+信息表(一业绩一张,克隆/裁剪)。

    返回 {"handled_tables": set[tbl元素], "detail_filled": int, "summary_filled": int}。
    没有任何记录(库未导入/未选派)时不动文档、返回空 handled——保持通用
    _fill_performance_table 的老行为。
    """
    result: dict[str, Any] = {"handled_tables": set(), "detail_filled": 0, "summary_filled": 0}
    pm_records = list(profile.get("similar_projects_pm") or [])
    td_records = list(profile.get("similar_projects_td") or [])
    selected = list(profile.get("selected_performance") or [])
    all_records = list(profile.get("similar_projects_all") or [])
    # 四个数据源全无货才整体不干活(老项目/库未导入,保持通用填表老行为)。
    # 注意不能只看 pm/td:投标人节靠 selected_performance+similar_projects_all 填,
    # 用户没选派经理时 pm/td 全空——旧闸在这直接退,导致选了业绩的表也一格不填(埇桥实测)。
    if not pm_records and not td_records and not selected and not all_records:
        return result

    body = document.element.body
    role: str | None = None
    qualifier = ""
    # 先一遍扫描:给每张相关表定归属 + 分类(明细/汇总),按文档流顺序。
    # 节键=(归属人, 评审限定词):同归属的"资格审查"和"详细评审"是两个节,
    # 各自要填全套记录,绝不能并组拆分。
    detail_units: list[dict[str, Any]] = []  # {key, tbl_el}
    summary_tables: list[dict[str, Any]] = []
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(child.itertext()).strip()
            switched = _role_of_title(text)
            if switched:
                q = _qualifier_of_title(text)
                if switched != role:
                    role, qualifier = switched, q
                elif q:
                    qualifier = q  # 同归属换限定词(资格审查→详细评审)=新节
            continue
        if child.tag != qn("w:tbl") or role is None:
            continue
        table = Table(child, document._body)
        if _is_similar_project_detail_table(table):
            detail_units.append({"key": (role, qualifier), "tbl_el": child})
        elif _is_summary_table(table):
            summary_tables.append({"key": (role, qualifier), "tbl_el": child})

    # ── 汇总情况表:按归属人填名单(+扩行) ──
    for entry in summary_tables:
        records = _records_for_role(entry["key"][0], profile)
        table = Table(entry["tbl_el"], document._body)
        result["handled_tables"].add(entry["tbl_el"])
        names = [str(r.get("project_name") or "").strip() for r in records]
        names = [n for n in names if n]
        if names:
            rows_before = len(table.rows)
            result["summary_filled"] += _fill_summary_table(table, names)
            if len(table.rows) > rows_before:
                # 表被扩行撑长 → 吸收其后的垫页空段,防整页空白
                note = _adjacent_note_p(entry["tbl_el"])
                tail = note if note is not None else entry["tbl_el"]
                _absorb_page_padding(tail)

    # ── 详细信息表:同节键(归属人+限定词)连续成组,组内 填→克隆补齐→裁剪多余 ──
    groups: list[dict[str, Any]] = []
    for unit in detail_units:
        if groups and groups[-1]["key"] == unit["key"]:
            groups[-1]["tbl_els"].append(unit["tbl_el"])
        else:
            groups.append({"key": unit["key"], "tbl_els": [unit["tbl_el"]]})

    for group in groups:
        records = _records_for_role(group["key"][0], profile)
        tbl_els: list[Any] = group["tbl_els"]
        for el in tbl_els:
            result["handled_tables"].add(el)
        if not records:
            continue  # 留白(总工匹配不到等),表已进 handled,不会被通用填表器乱填
        # 模板表不够 → 以最后一张为样板克隆(连同标题段/注段),前面加分页符
        template_el = tbl_els[-1]
        title_p = _adjacent_title_p(template_el)
        note_p = _adjacent_note_p(template_el)
        anchor = note_p if note_p is not None else template_el
        while len(tbl_els) < len(records):
            pieces: list[Any] = []
            if title_p is not None:
                title_copy = deepcopy(title_p)
                _strip_sectpr(title_copy)
                _ensure_page_break_before(title_copy)  # 一表一页,不靠独立分页空段
                pieces.append(title_copy)
            else:
                pieces.append(_page_break_paragraph())
            new_tbl = deepcopy(template_el)
            pieces.append(new_tbl)
            if note_p is not None:
                note_copy = deepcopy(note_p)
                _strip_sectpr(note_copy)
                pieces.append(note_copy)
            for piece in pieces:
                anchor.addnext(piece)
                anchor = piece
            tbl_els.append(new_tbl)
            result["handled_tables"].add(new_tbl)
        # 模板表多于记录 → 裁掉多余的**空**表(连同其标题段/注段),别留一堆空表占版面。
        # 值格有预印内容的表不删(那是招标给的示例/别人的地盘);带节属性的段也不删。
        while len(tbl_els) > len(records):
            extra = tbl_els[-1]
            if not _detail_table_is_blank(extra, document):
                break
            tbl_els.pop()
            for el in (_adjacent_title_p(extra), _adjacent_note_p(extra), extra):
                if (
                    el is not None
                    and el.getparent() is not None
                    and el.find(f".//{qn('w:sectPr')}") is None
                ):
                    el.getparent().remove(el)
        for tbl_el, record in zip(tbl_els, records):
            result["detail_filled"] += _fill_detail_table(
                Table(tbl_el, document._body), record
            )
        # 克隆/填值都会撑长这一组的版面 → 吸收组尾的垫页空段,防整页空白
        note = _adjacent_note_p(tbl_els[-1])
        tail = note if note is not None else tbl_els[-1]
        _absorb_page_padding(tail)

    logger.info(
        "类似项目自动填表:明细格 %d 个,汇总行 %d 行,经手表格 %d 张 (经理记录 %d 条,总工记录 %d 条)",
        result["detail_filled"], result["summary_filled"],
        len(result["handled_tables"]), len(pm_records), len(td_records),
    )
    return result
